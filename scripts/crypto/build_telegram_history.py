"""Recover source Telegram rows and accepted cash evidence, never legacy basis."""
import argparse
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
import hashlib
import json
from pathlib import Path

import openpyxl

from build_history_inventory import ROOT, git

BASE = 'crypto-task/crypto-reconciliation/work/'
CLASSIFICATION = BASE + 'p2p-color-audit/classification.json'
LOCK = 'crypto-task/chronological-rebuild/blocks/0001-0100/accepted/v1/cost-flow-0031-sources.json'
TELEGRAM = 'crypto-reconciliation/work/ton-cost-2024/inputs/telegram-history-converted.xlsx'
P2P = 'local-data/p2p_purchases_with_usd_rub_rates.backup_before_ton_usdt_equiv.xlsx'


def build(task_root, inventory):
    parent = json.loads((inventory / 'summary.json').read_text())
    commit = parent['source_commit']
    inputs = []

    def read_git(path):
        raw = git('show', commit + ':' + path)
        inputs.append({'path': path, 'sha256': hashlib.sha256(raw).hexdigest()})
        return json.loads(raw)

    lock = read_git(LOCK)
    classification = read_git(CLASSIFICATION)
    if inputs[-1]['sha256'] != lock['crypto-reconciliation/work/p2p-color-audit/classification.json']:
        raise ValueError('Classification differs from accepted evidence')

    def workbook(relative):
        path = task_root / relative
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest != lock[relative]:
            raise ValueError(f'Workbook differs from accepted evidence: {relative}')
        inputs.append({'path': str(path.resolve()), 'sha256': digest})
        return openpyxl.load_workbook(path, read_only=True, data_only=False)

    tg, payments = workbook(TELEGRAM), workbook(P2P)
    rows = []
    for number, values in enumerate(tg.active.iter_rows(values_only=True), 1):
        if number == 1:
            continue
        day, time, operation, direction, quantity, symbol, recipient, status, screenshot, note = values
        if not isinstance(day, datetime) or quantity is None or str(quantity).startswith('='):
            raise ValueError(f'Invalid Telegram source row {number}')
        rows.append({'source_key': f'telegram:row:{number}', 'source_row': number,
                     'date': day.date().isoformat(), 'time_as_in_source': time,
                     'operation': operation, 'direction': direction,
                     'quantity': str(Decimal(str(quantity))), 'source_asset': symbol,
                     'asset': 'TON' if symbol == 'GRAM' else symbol,
                     'recipient_as_in_source': recipient, 'source_status': status,
                     'screenshot_reference': screenshot, 'note': note,
                     'posting_status': 'not_mapped_not_posted'})
    indexed = {r['source_row']: r for r in rows}
    buys, excluded = [], []
    payment_rows = list(payments['Операции'].iter_rows(values_only=True))
    for entry in classification['rows']:
        if entry['venue'] == 'Bybit':
            excluded.append({'p2p_row': entry['row'], 'reason': 'use_bybit_export_do_not_import_twice'})
            continue
        if entry['venue'] != 'Telegram' or len(entry['telegram_rows']) != 1:
            raise ValueError('Unexpected purchase classification')
        number = entry['telegram_rows'][0]
        row = indexed[number]
        raw = payment_rows[entry['row'] - 1]
        raw_rub = Decimal(str(raw[13]))
        rub = raw_rub.quantize(Decimal('.01'))
        if abs(raw_rub - rub) > Decimal('0.00000001'):
            raise ValueError('Source payment has material sub-kopeck precision')
        if row['operation'] != 'Покупка: P2P' or row['date'] != entry['date'] or row['asset'] != entry['asset']:
            raise ValueError('Purchase link does not match Telegram source')
        if abs(Decimal(row['quantity']) - Decimal(entry['quantity'])) > Decimal('0.00000001'):
            raise ValueError('Purchase quantity differs from accepted evidence')
        if abs(rub - Decimal(entry['rub'])) > Decimal('0.00000001') or rub <= 0:
            raise ValueError(f'Purchase RUB differs at {entry["rub_cell"]}: workbook={rub}, accepted={entry["rub"]}')
        if entry['status'] not in {'owner_actual', 'owner_allowed_estimate'}:
            raise ValueError('Unsupported source value quality')
        buys.append({**row, 'kind': 'buy_fiat', 'fiat_currency': 'RUB', 'fiat_amount': str(rub),
                     'basis_quality': 'known' if entry['status'] == 'owner_actual' else 'estimated',
                     'amount_evidence': {'classification': CLASSIFICATION, 'status': entry['status'],
                                         'workbook': P2P, 'cell': entry['rub_cell'], 'stored_numeric_value': str(raw_rub)},
                     'posting_status': 'source_ready_accounts_order_and_quality_mapping_required'})
    if {r['source_row'] for r in buys} != {r['source_row'] for r in rows if r['operation'] == 'Покупка: P2P'} or len(buys) != len({r['source_row'] for r in buys}):
        raise ValueError('Telegram purchase coverage mismatch')

    prefix = 'crypto-task/chronological-rebuild/blocks/'
    not_path = prefix + '0101-0200/accepted/v2/owner-NOT-exchange.json'
    dogs_path = prefix + '0101-0200/accepted/v2/owner-DOGS-exchange.json'
    later_path = prefix + '0301-0400/accepted/v2/owner-telegram-clarification.json'
    not_answer, dogs_answer, later = read_git(not_path), read_git(dogs_path), read_git(later_path)
    patches = [
        ('2024-06-12', 'NOT', not_answer['quantity_received'], not_answer['spent_asset'], not_answer['quantity_spent'], not_path),
        ('2024-08-28', 'DOGS', dogs_answer['quantity_received'], dogs_answer['asset_spent'], dogs_answer['quantity_spent'], dogs_path),
        ('2024-11-22', 'USDT', later['November_USDT_received'], 'TON', later['November_TON_spent'], later_path),
        ('2024-12-21', 'TON', later['December_TON_received'], 'USDT', later['December_USDT_spent'], later_path),
    ]
    swaps = []
    for day, asset, qty, spent_asset, spent_qty, evidence in patches:
        matches = [r for r in rows if r['date'] == day and r['asset'] == asset
                   and Decimal(r['quantity']) == Decimal(qty) and r['operation'].startswith('Обмен')]
        if len(matches) != 1:
            raise ValueError('Owner swap clarification has no unique source row')
        swaps.append({**matches[0], 'spent_asset': spent_asset, 'spent_quantity': spent_qty,
                      'owner_evidence': evidence, 'historical_rub_value': None,
                      'valuation_status': 'unknown_not_legacy_carried_basis'})

    precision_path = BASE + 'telegram-usdt-precision/result.json'
    precision = read_git(precision_path)
    events = {r['event_id']: r for r in json.loads((inventory / 'main-events.json').read_text())}
    received = defaultdict(Decimal)
    for movement in json.loads((inventory / 'main-movements.json').read_text()):
        if movement['asset'] == 'USDT' and movement['quantity_delta'] is not None:
            quantity = Decimal(movement['quantity_delta'])
            if quantity > 0:
                received[movement['event_id']] += quantity
    links = []
    for item in precision['matches']:
        row = indexed[item['telegram_row']]
        event = events.get(item['event'])
        if event is None or row['asset'] != 'USDT' or row['operation'] != 'Вывод' or Decimal(row['quantity']) != Decimal(item['document_amount']):
            raise ValueError('USDT source bridge does not match accepted main event')
        if received[item['event']] != Decimal(item['chain_received']):
            raise ValueError('Saved bridge quantity differs from accepted main ledger')
        links.append({'telegram_row': row['source_row'], 'event_id': event['event_id'],
                      'event_no': event['event_no'], 'document_quantity': row['quantity'],
                      'chain_received': item['chain_received'], 'source': precision_path,
                      'status': 'saved_quantity_evidence_not_new_fiat_purchase'})
    tg.close()
    payments.close()
    quality = Counter(r['basis_quality'] for r in buys)
    totals = {q: str(sum(Decimal(r['fiat_amount']) for r in buys if r['basis_quality'] == q)) for q in quality}
    summary = {'source_commit': commit, 'source_rows': len(rows), 'purchase_rows': len(buys),
               'purchase_quality_counts': dict(quality), 'rub_by_quality': totals,
               'excluded_duplicate_bybit_rows': len(excluded), 'owner_swap_clarifications': len(swaps),
               'usdt_links_to_main': len(links), 'unknown_source_times': sum(r['time_as_in_source'] is None for r in rows),
               'status': 'source_prepared_not_posted', 'inputs': inputs}
    return {'summary.json': summary, 'telegram-source-rows.json': rows, 'telegram-purchases.json': buys,
            'telegram-owner-swaps.json': swaps, 'telegram-main-usdt-links.json': links,
            'excluded-duplicate-bybit.json': excluded}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--task-root', type=Path, required=True)
    parser.add_argument('--inventory', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    args = parser.parse_args()
    out = args.output_dir.resolve()
    if out.exists() or not out.is_relative_to((ROOT / 'outputs').resolve()):
        parser.error('Use a new directory under ignored outputs/')
    result = build(args.task_root, args.inventory)
    out.mkdir(parents=True)
    for name, data in result.items():
        (out / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({k: v for k, v in result['summary.json'].items() if k != 'inputs'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
